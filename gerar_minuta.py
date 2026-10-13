import os
import re
import sys
import copy
import glob
import time
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from num2words import num2words


# =============================================================================
# CONFIGURACAO
#
# A raiz dos arquivos de decreto (BASE_USUARIO) NAO fica fixa no codigo, para
# funcionar em qualquer maquina. A ordem de resolucao e:
#
#   1) Variavel de ambiente MINUTA_BASE, se definida. Ex.:
#        Windows : setx MINUTA_BASE "C:\Users\SEU_ID\OneDrive - CAMG\...\Decretos"
#        WSL/Linux: export MINUTA_BASE="/mnt/c/Users/SEU_ID/.../Decretos"
#   2) Padrao: <pasta do usuario logado>\OneDrive - CAMG\General\@dcmefo\2026\Decretos
#      (Path.home() resolve o C:\Users\<seu_id> automaticamente)
# =============================================================================
_DEFAULT_BASE = Path.home() / "OneDrive - CAMG" / "General" / "@dcmefo" / "2026" / "Decretos"
BASE_USUARIO = os.environ.get("MINUTA_BASE", str(_DEFAULT_BASE))

# Pasta raiz das minutas (a mesma usada pela conferencia e pela geracao)
BASE_MINUTAS = os.path.join(BASE_USUARIO, "Minutas Finalizadas")

# Template da minuta
TEMPLATE_PATH = os.path.join(
    BASE_USUARIO,
    "Manual decreto", "minutas padrão", "minuta xxx.docx",
)


# #############################################################################
# PARTE 1 — CONFERENCIA (gera o conferencia.xlsx)
# #############################################################################

# =============================================================================
# EXTENSO EM PORTUGUES (pt-BR) — usado nos incisos do conferencia.xlsx
# =============================================================================
UNI = ["", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove",
       "dez", "onze", "doze", "treze", "catorze", "quinze", "dezesseis",
       "dezessete", "dezoito", "dezenove"]
DEZ = ["", "", "vinte", "trinta", "quarenta", "cinquenta",
       "sessenta", "setenta", "oitenta", "noventa"]
CEN = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos",
       "seiscentos", "setecentos", "oitocentos", "novecentos"]


def _ate_999(n: int) -> str:
    if n == 0:
        return ""
    if n == 100:
        return "cem"
    c, r = divmod(n, 100)
    partes = []
    if c:
        partes.append(CEN[c])
    if r:
        if r < 20:
            partes.append(UNI[r])
        else:
            d, u = divmod(r, 10)
            partes.append(DEZ[d] if u == 0 else f"{DEZ[d]} e {UNI[u]}")
    return " e ".join(partes)


def _inteiro_extenso(n: int) -> str:
    if n == 0:
        return "zero"
    bilhoes, r = divmod(n, 1_000_000_000)
    milhoes, r = divmod(r, 1_000_000)
    milhares, centenas = divmod(r, 1_000)
    partes = []
    if bilhoes:
        partes.append(f"{_ate_999(bilhoes)} {'bilhão' if bilhoes == 1 else 'bilhões'}")
    if milhoes:
        partes.append(f"{_ate_999(milhoes)} {'milhão' if milhoes == 1 else 'milhões'}")
    if milhares:
        partes.append("mil" if milhares == 1 else f"{_ate_999(milhares)} mil")
    if centenas:
        partes.append(_ate_999(centenas))
    # entre blocos (bilhao/milhao/mil/centenas) usa-se apenas espaco,
    # seguindo o padrao do arquivo de referencia
    return " ".join(partes)


def valor_extenso_incisos(valor: float) -> str:
    """Extenso proprio da conferencia (usado no texto dos incisos do xlsx)."""
    d = Decimal(str(valor)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    inteiro = int(d)
    centavos = int((d - inteiro) * 100)
    partes = []
    if inteiro:
        partes.append(f"{_inteiro_extenso(inteiro)} {'real' if inteiro == 1 else 'reais'}")
    if centavos:
        nome = "centavo" if centavos == 1 else "centavos"
        ext = f"{_inteiro_extenso(centavos)} {nome}"
        partes.append(f"e {ext}" if partes else ext)
    return " ".join(partes) if partes else "zero reais"


def fmt_valor(valor: float) -> str:
    """Formata 22224054.11 -> '22.224.054,11'."""
    d = Decimal(str(valor)).quantize(Decimal("0.01"), ROUND_HALF_UP)
    s = f"{d:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def romano(n: int) -> str:
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"),
            (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


# =============================================================================
# LEITURA DOS .DOCX
# =============================================================================
def ler_paragrafos_docx(caminho: str) -> list[str]:
    """Retorna todos os paragrafos de texto de um .docx (corpo + tabelas)."""
    doc = Document(caminho)
    textos = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for p in celula.paragraphs:
                    textos.append(p.text)
    return textos


# =============================================================================
# EXTRACAO E CONSOLIDACAO DOS INCISOS
# =============================================================================
PADRAO_FRASE = re.compile(r"(?i)(do\s+saldo|do\s+excesso)")
PADRAO_NO_VALOR = re.compile(r"(?i)no\s+valor")
# captura tudo de "do saldo" / "do excesso" ate (sem incluir) "no valor de"
PADRAO_CHAVE = re.compile(
    r"(?i)(do\s+(?:saldo|excesso).*?)\s*,?\s*no\s+valor\s+de"
)
# captura o valor monetario logo apos "R$"
PADRAO_VALOR = re.compile(r"R\$\s*(\d[\d\.]*,\d{2})")


def parse_valor_br(s: str) -> float:
    """'16.380.103,30' -> 16380103.30"""
    return float(s.replace(".", "").replace(",", "."))


def processar_pasta(pasta: str) -> list[dict]:
    """Le todos os .docx da pasta e retorna lista de dicts {chave, valor}."""
    arquivos = sorted(glob.glob(os.path.join(pasta, "*.docx")))
    arquivos = [a for a in arquivos if not os.path.basename(a).startswith("~$")]

    if not arquivos:
        raise FileNotFoundError(f"Nenhum arquivo .docx encontrado em: {pasta}")

    registros = []
    for caminho in arquivos:
        try:
            paragrafos = ler_paragrafos_docx(caminho)
        except Exception as e:
            print(f"Erro ao ler {caminho}: {e}")
            continue

        for texto in paragrafos:
            # remove espaco entre "R$" e o numero (R$ 16.380 -> R$16.380)
            t = re.sub(r"(R\$)\s+", r"\1", texto)
            # so paragrafos que contenham as palavras-chave E "no valor"
            if not (PADRAO_FRASE.search(t) and PADRAO_NO_VALOR.search(t)):
                continue
            m_chave = PADRAO_CHAVE.search(t)
            m_valor = PADRAO_VALOR.search(t)
            if not m_chave or not m_valor:
                continue
            chave = m_chave.group(1).strip()
            valor = parse_valor_br(m_valor.group(1))
            registros.append({"chave": chave, "valor": valor})

    return registros


def consolidar(registros: list[dict]) -> list[dict]:
    """Agrupa por chave normalizada (case-insensitive + espacos normalizados),
    soma os valores e preserva a primeira capitalizacao encontrada."""
    grupos: dict[str, dict] = {}
    ordem: list[str] = []
    for r in registros:
        chave_norm = re.sub(r"\s+", " ", r["chave"].lower()).strip()
        if chave_norm not in grupos:
            grupos[chave_norm] = {"chave_texto": r["chave"], "valor": 0.0}
            ordem.append(chave_norm)
        grupos[chave_norm]["valor"] += r["valor"]

    saida = []
    for i, chave_norm in enumerate(ordem):
        g = grupos[chave_norm]
        num = romano(i + 2)  # comeca em II como no script original
        valor = round(g["valor"], 2)
        inciso = (
            f"{num} - {g['chave_texto']} no valor de "
            f"R${fmt_valor(valor)} ({valor_extenso_incisos(valor)});"
        )
        saida.append({"incisos": inciso, "valor": valor})
    return saida


# =============================================================================
# ESCRITA DO EXCEL
# =============================================================================
def salvar_xlsx(linhas: list[dict], caminho_saida: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "base"
    ws.append(["incisos", "valor"])
    for r in linhas:
        ws.append([r["incisos"], r["valor"]])
    # largura util
    ws.column_dimensions["A"].width = 120
    ws.column_dimensions["B"].width = 18
    wb.save(caminho_saida)


# =============================================================================
# ROTINA DA CONFERENCIA (antes: main() do script de conferencia)
# =============================================================================
def rodar_conferencia(pasta: str) -> None:
    """Roda a conferencia sobre a pasta informada e grava conferencia.xlsx."""
    registros = processar_pasta(pasta)
    saida = os.path.join(pasta, "conferencia.xlsx")

    if not registros:
        # Mesmo sem incisos de 'do saldo'/'do excesso', cria o Excel
        # apenas com o cabecalho (colunas 'incisos' e 'valor'), sem conteudo.
        salvar_xlsx([], saida)
        print("Nenhum inciso de 'do saldo'/'do excesso' com 'no valor de' encontrado.")
        print(f"Arquivo criado vazio (apenas cabeçalho): {saida}")
        return

    linhas = consolidar(registros)
    salvar_xlsx(linhas, saida)
    print(f"Processo concluído. {len(registros)} incisos lidos, "
          f"{len(linhas)} após consolidação. Arquivo: {saida}")


# #############################################################################
# PARTE 2 — GERACAO DA MINUTA (le o conferencia.xlsx gerado acima)
# #############################################################################

# ---------------------------------------------------------------------------
# 1) Conversoes de acentuacao aplicadas aos titulos de orgaos / cabecalhos
# ---------------------------------------------------------------------------
WORD_FIXES = {
    "POLICIA": "POLÍCIA",
    "PREVIDENCIA": "PREVIDÊNCIA",
    "EDUCACAO": "EDUCAÇÃO",
    "CONTINGENCIA": "CONTINGÊNCIA",
    "SUPLEMENTACAO": "SUPLEMENTAÇÃO",
    "DOTACOES": "DOTAÇÕES",
    "ORCAMENTARIAS": "ORÇAMENTÁRIAS",
    "ANULACAO": "ANULAÇÃO",
    "SAUDE": "SAÚDE",
    "PECUARIA": "PECUÁRIA",
    "FUNDACAO": "FUNDAÇÃO",
    "JUSTICA": "JUSTIÇA",
    "SEGURANCA": "SEGURANÇA",
    "JOAO": "JOÃO",
    "PATRIMONIO": "PATRIMÔNIO",
    "HISTORICO": "HISTÓRICO",
    "ARTISTICO": "ARTÍSTICO",
    "AGROPECUARIA": "AGROPECUÁRIA",
}


def reacentuar(linha: str) -> str:
    """Re-acentua palavras-chave de um titulo/cabecalho."""
    def sub(match):
        w = match.group(0)
        return WORD_FIXES.get(w, w)
    return re.sub(r"\b\w+\b", sub, linha)


# ---------------------------------------------------------------------------
# 2) Conversor numerico -> extenso em portugues (moeda BRL) via num2words
# ---------------------------------------------------------------------------
def valor_extenso(valor: float) -> str:
    texto = num2words(valor, lang="pt_BR", to="currency")
    return texto.replace(",", "")


def fmt_brl(valor: float) -> str:
    inteiro, _, dec = f"{valor:.2f}".partition(".")
    inteiro_fmt = ""
    for i, ch in enumerate(reversed(inteiro)):
        if i and i % 3 == 0:
            inteiro_fmt = "." + inteiro_fmt
        inteiro_fmt = ch + inteiro_fmt
    return f"{inteiro_fmt},{dec}"


# ---------------------------------------------------------------------------
# 3) Parser do arquivo Minuta00NN.TXT
# ---------------------------------------------------------------------------
RE_LINHA_DOTACAO = re.compile(r"^\s*(\d{4}\.[\d\-\.]+)\s+([\d\.\,]+)\s*$")
RE_NUMERO_SIAFI = re.compile(r"Sob\s+o\s+n[úu]mero\s+(\d+)", re.IGNORECASE)


def parse_txt(caminho_txt: str) -> dict:
    with open(caminho_txt, encoding="latin-1") as f:
        linhas = f.readlines()

    resultado = {
        "numero": None,
        "suplementacao": [],
        "anulacao": [],
        "total_supl": None,
        "total_anul": None,
    }

    secao = None
    org_atual = None
    secao_atual = None

    for linha_bruta in linhas:
        linha = linha_bruta.rstrip("\n").strip()
        if not linha:
            continue

        m = RE_NUMERO_SIAFI.search(linha)
        if m and not resultado["numero"]:
            resultado["numero"] = m.group(1)

        if linha.startswith("ANEXO AO DECRETO"):
            continue
        if linha.startswith("(Registrado no SIAFI"):
            continue

        # Detecta inicio da secao de suplementacao (singular ou plural)
        if re.match(r"SUPLEMENTACAO\s+DA[S]?\s+SEGUINTE[S]?", linha):
            secao = "supl"
            secao_atual = resultado["suplementacao"]
            org_atual = None
            continue
        # Descarta continuacoes do cabecalho de suplementacao e anulacao
        if re.match(r"SE\s+REFERE\s+O\s+ART[\.\s]", linha):
            continue
        if re.match(r"2[,\s]", linha) and "DESTE DECRETO" in linha:
            continue
        # Detecta inicio da secao de anulacao (singular ou plural)
        if re.match(r"ANULACAO\s+DA[S]?\s+SEGUINTE[S]?", linha):
            secao = "anul"
            secao_atual = resultado["anulacao"]
            org_atual = None
            continue

        if linha.startswith("TOTAL DA SUPLEMENTACAO"):
            resultado["total_supl"] = linha.split()[-1]
            continue
        if linha.startswith("TOTAL DA ANULACAO"):
            resultado["total_anul"] = linha.split()[-1]
            continue

        if secao is None:
            continue

        m = RE_LINHA_DOTACAO.match(linha_bruta.rstrip("\n"))
        if m:
            if org_atual is None:
                org_atual = ("", [])
                secao_atual.append(org_atual)
            org_atual[1].append(linha_bruta.rstrip("\n").strip())
            continue

        if linha == "R$":
            continue

        titulo = reacentuar(linha)
        org_atual = (titulo, [])
        secao_atual.append(org_atual)

    return resultado


# ---------------------------------------------------------------------------
# 3.1) Verificacao de fontes e IPU especiais
#
# Regras:
#   - IPU == 0 (qualquer fonte)  -> "Minuta contém Recursos do PROPAG"
#   - Fonte == 89 (qualquer IPU) -> "Minuta contém Recursos do PROPAG"
#   - Fonte == 95 (qualquer IPU) -> "Minuta contém Recursos de Brumadinho"
#   - Fonte == 80 (qualquer IPU) -> "Minuta contém Recursos de Mariana"
#
# O campo fonte-ipu fica no ultimo segmento da dotacao, apos o ultimo '-':
#   ex: 1401.06182052-4.115-0001-3390-0-70.1  ->  "70.1"  -> fonte=70, ipu=1
# ---------------------------------------------------------------------------

# Regex que captura o trecho fonte.ipu no final da linha de dotacao
RE_FONTE_IPU = re.compile(
    r"\d{4}\.[\d\-\.]+-(\d+)\.(\d+)\s+[\d\.,]+\s*$"
)


def verificar_alertas(caminho_txt: str) -> list[tuple[str, list[str]]]:
    """
    Lê o TXT e retorna lista de (mensagem, [linhas_que_causaram]).
    Cada mensagem aparece no maximo uma vez, com todas as linhas associadas.
    """
    alertas: dict[str, list[str]] = {}

    def registrar(msg: str, linha: str):
        alertas.setdefault(msg, []).append(linha.strip())

    with open(caminho_txt, encoding="latin-1") as f:
        for linha_bruta in f:
            linha = linha_bruta.rstrip("\n")
            m = RE_FONTE_IPU.search(linha)
            if not m:
                continue

            fonte = int(m.group(1))
            ipu   = int(m.group(2))

            if ipu == 0 or fonte == 89:
                registrar("⚠️  Minuta contém Recursos do PROPAG", linha)
            if fonte == 95:
                registrar("⚠️  Minuta contém Recursos de Brumadinho", linha)
            if fonte == 80:
                registrar("⚠️  Minuta contém Recursos de Mariana", linha)

    return list(alertas.items())


def exibir_alertas_e_confirmar(caminho_txt: str) -> bool:
    """
    Exibe os alertas encontrados no TXT.
    Retorna True se deve prosseguir, False se o usuario cancelou.
    Se nao houver alertas, retorna True silenciosamente.
    """
    alertas = verificar_alertas(caminho_txt)

    if not alertas:
        return True

    print()
    print("=" * 60)
    print("  ATENÇÃO — VERIFICAÇÃO DA MINUTA")
    print("=" * 60)
    for msg, linhas in alertas:
        print(f"\n  {msg}")
        for l in linhas:
            print(f"    → {l}")
    print()
    print("=" * 60)

    while True:
        resp = input("  Deseja continuar mesmo assim? [s/n]: ").strip().lower()
        if resp in ("s", "sim", "y", "yes"):
            print()
            return True
        if resp in ("n", "nao", "não", "no"):
            print("\n  Geração cancelada pelo usuário.")
            return False
        print("  Resposta inválida. Digite 's' para continuar ou 'n' para cancelar.")


# ---------------------------------------------------------------------------
# 3.2) Verificacao de balanco: Suplementacao == Anulacao + Incisos
# ---------------------------------------------------------------------------
def _parse_brl(valor_str: str) -> float:
    """Converte string BRL '1.234.567,89' para float."""
    return float(valor_str.replace(".", "").replace(",", "."))


def verificar_balanco(caminho_txt: str, caminho_xlsx: str) -> bool:
    """
    Compara Suplementacao x (Anulacao + Incisos).
    Exibe aviso e pergunta se deve continuar quando ha divergencia.
    Retorna True para prosseguir, False para cancelar.
    """
    dados = parse_txt(caminho_txt)
    incisos = parse_conferencia(caminho_xlsx)

    val_supl  = _parse_brl(dados["total_supl"])
    val_anul  = _parse_brl(dados["total_anul"])
    val_incis = sum(v for _, v in incisos)
    val_art2  = val_anul + val_incis

    # Sem tolerancia: os valores devem bater exatamente ate o centavo.
    # A comparacao e feita em centavos inteiros apenas para descartar ruido
    # de representacao de ponto flutuante (ex.: 1,15 * 100 = 114,9999...),
    # e NAO para admitir qualquer diferenca real de centavos.
    if round(val_supl * 100) == round(val_art2 * 100):
        return True

    diff = val_art2 - val_supl
    fmt_supl  = fmt_brl(val_supl)
    fmt_art2  = fmt_brl(val_art2)
    fmt_diff  = fmt_brl(abs(diff))

    print()
    print("=" * 60)
    print("  ATENÇÃO — DIVERGÊNCIA DE VALORES")
    print("=" * 60)
    print(f"\n  Valor Suplementado (Art. 1º) : R$ {fmt_supl}")
    print(f"  Anulação (TXT)               : R$ {fmt_brl(val_anul)}")
    if incisos:
        print(f"  Incisos (conferencia.xlsx)   : R$ {fmt_brl(val_incis)}")
    print(f"  Total Art. 2º                : R$ {fmt_art2}")
    print()

    if diff > 0:
        print(f"  ⚠️  Valor do Art. 2º é MAIOR que o Valor Suplementado")
        print(f"      (excesso de R$ {fmt_diff})")
    else:
        print(f"  ⚠️  Valor do Art. 2º é MENOR que o Valor Suplementado")
        print(f"      (falta de R$ {fmt_diff})")

    print()
    print("=" * 60)

    while True:
        resp = input("  Deseja continuar mesmo assim? [s/n]: ").strip().lower()
        if resp in ("s", "sim", "y", "yes"):
            print()
            return True
        if resp in ("n", "nao", "não", "no"):
            print("\n  Geração cancelada pelo usuário.")
            return False
        print("  Resposta inválida. Digite 's' para continuar ou 'n' para cancelar.")


# ---------------------------------------------------------------------------
# 4) Parser do conferencia.xlsx
# ---------------------------------------------------------------------------
def parse_conferencia(caminho_xlsx: str):
    wb = load_workbook(caminho_xlsx, data_only=True)
    ws = wb.active
    incisos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        texto, valor = row[0], row[1]
        if texto is None:
            continue
        incisos.append((str(texto).strip(), float(valor) if valor is not None else 0.0))
    return incisos


# ---------------------------------------------------------------------------
# 5) Construcao do documento .docx
# ---------------------------------------------------------------------------
def substituir_em_runs(paragrafo, mapa: dict):
    texto_inteiro = "".join(r.text for r in paragrafo.runs)
    novo_texto = texto_inteiro
    houve_troca = False
    for chave, valor in mapa.items():
        if chave in novo_texto:
            novo_texto = novo_texto.replace(chave, valor)
            houve_troca = True
    if not houve_troca:
        return
    if paragrafo.runs:
        paragrafo.runs[0].text = novo_texto
        for r in paragrafo.runs[1:]:
            r.text = ""


def clonar_paragrafo_apos(referencia, texto: str, alinhamento=None):
    novo_xml = copy.deepcopy(referencia._p)
    referencia._p.addnext(novo_xml)
    novo = Paragraph(novo_xml, referencia._parent)
    for r in list(novo.runs)[1:]:
        r._r.getparent().remove(r._r)
    if novo.runs:
        novo.runs[0].text = texto
    else:
        novo.add_run(texto)
    if alinhamento is not None:
        novo.alignment = alinhamento
    return novo


def adicionar_paragrafos_sequencia(referencia, textos):
    atual = referencia
    novos = []
    for txt in textos:
        atual = clonar_paragrafo_apos(atual, txt)
        novos.append(atual)
    return novos


def gerar_minuta(template_path: str, txt_path: str, xlsx_path: str, saida_path: str):
    dados = parse_txt(txt_path)
    incisos_extra = parse_conferencia(xlsx_path)

    total_supl_str = dados["total_supl"]
    total_anul_str = dados["total_anul"]
    total_supl_val = float(total_supl_str.replace(".", "").replace(",", "."))
    total_anul_val = float(total_anul_str.replace(".", "").replace(",", "."))

    numero = dados["numero"] or ""

    extenso_supl = valor_extenso(total_supl_val)
    extenso_anul = valor_extenso(total_anul_val)

    doc = Document(template_path)
    paragrafos = doc.paragraphs

    substituir_em_runs(paragrafos[1], {"xxx": total_supl_str})

    p4 = paragrafos[4]
    texto_p4 = "".join(r.text for r in p4.runs)
    texto_p4 = texto_p4.replace("xxx por extenso", extenso_supl, 1)
    texto_p4 = texto_p4.replace("xxx", total_supl_str, 1)
    if p4.runs:
        p4.runs[0].text = texto_p4
        for r in p4.runs[1:]:
            r.text = ""

    p5 = paragrafos[5]   # "Art. 2º – Para atender... serão utilizados recursos provenientes:"
    p6 = paragrafos[6]   # "I – da anulação... R$xxx (xxx por extenso);"
    p7 = paragrafos[7]   # "Xxx" (placeholder de incisos extras)

    if incisos_extra:
        # Caso com incisos: comportamento original — preenche inciso I e adiciona os demais
        texto_p6 = "".join(r.text for r in p6.runs)
        texto_p6 = texto_p6.replace("xxx por extenso", extenso_anul, 1)
        texto_p6 = texto_p6.replace("xxx", total_anul_str, 1)
        if p6.runs:
            p6.runs[0].text = texto_p6
            for r in p6.runs[1:]:
                r.text = ""
    else:
        # Caso sem incisos: reescreve p5 como paragrafo unico do Art. 2º e apaga p6 e p7
        texto_art2 = (
            f"Art. 2\u00ba \u2013 Para atender ao disposto no art.1\u00ba ser\u00e1 utilizado "
            f"recurso proveniente da anula\u00e7\u00e3o das dota\u00e7\u00f5es or\u00e7ament\u00e1rias "
            f"indicadas no Anexo, no valor de R${total_anul_str} "
            f"({extenso_anul})."
        )
        if p5.runs:
            p5.runs[0].text = texto_art2
            for r in p5.runs[1:]:
                r.text = ""
        else:
            p5.add_run(texto_art2)
        # Remove fisicamente os paragrafos do inciso I e do placeholder Xxx
        for p_remover in (p6, p7):
            p_remover._element.getparent().remove(p_remover._element)

    if incisos_extra:
        textos_norm = []
        n = len(incisos_extra)
        for idx, (texto, _val) in enumerate(incisos_extra):
            t = texto.replace("- ", "– ", 1)
            ultimo = idx == n - 1
            if ultimo:
                if t.endswith(";"):
                    t = t[:-1] + "."
                elif not t.endswith("."):
                    t += "."
            else:
                if not t.endswith(";"):
                    if t.endswith("."):
                        t = t[:-1] + ";"
                    else:
                        t += ";"
            textos_norm.append(t)

        primeiro_texto_norm = textos_norm[0]
        if p7.runs:
            p7.runs[0].text = primeiro_texto_norm
            for r in p7.runs[1:]:
                r.text = ""
        else:
            p7.add_run(primeiro_texto_norm)

        ref = p7
        for texto_norm in textos_norm[1:]:
            ref = clonar_paragrafo_apos(ref, texto_norm)

    for p in doc.paragraphs:
        if "Sob o" in p.text and "xxx" in p.text:
            substituir_em_runs(p, {"xxx": numero})
            break

    p_titulo = None
    for p in doc.paragraphs:
        if p.text.startswith("SE REFERE O ART.1") and "xxx" in p.text:
            p_titulo = p
            break

    if p_titulo is None:
        raise RuntimeError("Nao encontrei o paragrafo de titulo da tabela.")

    primeiro_org = dados["suplementacao"][0][0] if dados["suplementacao"] else ""
    for r in p_titulo.runs:
        if "xxx" in r.text:
            r.text = r.text.replace("xxx", primeiro_org)
            break

    linhas_a_inserir = []

    linhas_a_inserir.append("R$")
    primeira_org_done = False
    for org, dotacoes in dados["suplementacao"]:
        if not primeira_org_done:
            primeira_org_done = True
        else:
            linhas_a_inserir.append(org)
        for d in dotacoes:
            linhas_a_inserir.append(d)

    linhas_a_inserir.append(
        f"TOTAL DA SUPLEMENTAÇÃO                                   {total_supl_str}"
    )

    linhas_a_inserir.append(
        "ANULAÇÃO DAS SEGUINTES DOTAÇÕES ORÇAMENTÁRIAS A QUE SE REFERE O ART.2, INCISO I, DESTE DECRETO:"
    )

    primeira_org_done = False
    for org, dotacoes in dados["anulacao"]:
        linhas_a_inserir.append(org)
        if not primeira_org_done:
            linhas_a_inserir.append("R$")
            primeira_org_done = True
        for d in dotacoes:
            linhas_a_inserir.append(d)

    linhas_a_inserir.append(
        f"TOTAL DA ANULAÇÃO                                        {total_anul_str}"
    )

    adicionar_paragrafos_sequencia(p_titulo, linhas_a_inserir)

    doc.save(saida_path)
    return saida_path


# ---------------------------------------------------------------------------
# 6) Localizacao dos arquivos da subpasta
# ---------------------------------------------------------------------------
def localizar_txt(pasta: str) -> str:
    """Procura um arquivo Minuta*.TXT (qualquer caixa) na pasta."""
    candidatos = []
    for nome in os.listdir(pasta):
        nome_lower = nome.lower()
        if nome_lower.startswith("minuta") and nome_lower.endswith(".txt"):
            candidatos.append(os.path.join(pasta, nome))

    if not candidatos:
        raise FileNotFoundError(
            f"Nenhum arquivo 'Minuta*.TXT' encontrado em:\n  {pasta}"
        )
    if len(candidatos) > 1:
        print("[AVISO] Mais de um TXT encontrado na pasta. Usarei o primeiro:")
        for c in candidatos:
            print(f"  - {c}")
    return candidatos[0]


def extrair_numero_do_txt(nome_arquivo: str) -> str:
    """Extrai o numero da minuta do nome do arquivo, ex: Minuta0083.TXT -> 83."""
    base = os.path.basename(nome_arquivo)
    m = re.search(r"(\d+)", base)
    if not m:
        return ""
    return str(int(m.group(1)))  # remove zeros a esquerda


# =============================================================================
# ROTINA DA GERACAO (antes: main() do script de geracao)
# =============================================================================
def rodar_minuta(pasta_minuta: str) -> bool:
    """
    Gera a minuta a partir da pasta informada.
    Retorna True se a minuta foi gerada, False se cancelada/erro.
    """
    # Validacoes
    if not os.path.exists(TEMPLATE_PATH):
        print(f"[ERRO] Template nao encontrado:\n  {TEMPLATE_PATH}", file=sys.stderr)
        return False

    if not os.path.isdir(pasta_minuta):
        print(f"[ERRO] Pasta da minuta nao encontrada:\n  {pasta_minuta}", file=sys.stderr)
        return False

    xlsx_path = os.path.join(pasta_minuta, "conferencia.xlsx")
    if not os.path.exists(xlsx_path):
        print(f"[ERRO] 'conferencia.xlsx' nao encontrado em:\n  {pasta_minuta}", file=sys.stderr)
        return False

    try:
        txt_path = localizar_txt(pasta_minuta)
    except FileNotFoundError as e:
        print(f"[ERRO] {e}", file=sys.stderr)
        return False

    numero = extrair_numero_do_txt(txt_path)
    nome_saida = f"minuta {numero}.docx" if numero else "minuta - gerada.docx"
    saida_path = os.path.join(pasta_minuta, nome_saida)

    print(f"\n  Template : {TEMPLATE_PATH}")
    print(f"  TXT      : {txt_path}")
    print(f"  XLSX     : {xlsx_path}")
    print(f"  Saida    : {saida_path}")

    # Verificacoes ANTES de gerar o .docx (so pedem confirmacao quando ha algo a revisar)
    if not exibir_alertas_e_confirmar(txt_path):
        return False

    if not verificar_balanco(txt_path, xlsx_path):
        return False

    try:
        gerar_minuta(TEMPLATE_PATH, txt_path, xlsx_path, saida_path)
    except Exception as e:
        print(f"\n[ERRO] Falha ao gerar minuta: {e}", file=sys.stderr)
        return False

    print(f"OK -> {saida_path}")
    return True


# #############################################################################
# PARTE 3 — ORQUESTRACAO (fluxo unico: conferencia -> confirma -> minuta)
# #############################################################################
def confirmar_prosseguimento() -> bool:
    """Pergunta se deve seguir para a geracao da minuta. True=s, False=n."""
    print()
    print("=" * 60)
    while True:
        resp = input("  Conferência concluída. Prosseguir com a geração da minuta? [s/n]: ").strip().lower()
        if resp in ("s", "sim", "y", "yes"):
            return True
        if resp in ("n", "nao", "não", "no"):
            return False
        print("  Resposta inválida. Digite 's' para gerar a minuta ou 'n' para encerrar.")


def main():
    print("=" * 60)
    print("  Gerador de Minuta de Decreto — Conferência + Geração")
    print("=" * 60)

    # Subpasta: argumento (sys.argv[1]) ou prompt — perguntada UMA unica vez
    if len(sys.argv) >= 2:
        subpasta = sys.argv[1]
    else:
        print(f"\nBase: {BASE_MINUTAS}")
        print("Informe a subpasta da minuta (ex: 05 maio\\minuta 83):")
        subpasta = input("> ").strip().strip('"').strip("'")

    if not subpasta:
        print("[ERRO] Subpasta não informada.", file=sys.stderr)
        sys.exit(1)

    pasta = os.path.join(BASE_MINUTAS, subpasta)
    if not os.path.isdir(pasta):
        print(f"[ERRO] Pasta da minuta não encontrada:\n  {pasta}", file=sys.stderr)
        sys.exit(1)

    # ---- Etapa 1: Conferencia ----
    print("\n" + "-" * 60)
    print("  Etapa 1 — Conferência")
    print("-" * 60)
    try:
        rodar_conferencia(pasta)
    except Exception as e:
        print(f"\n[ERRO] Falha na conferência: {e}", file=sys.stderr)
        sys.exit(1)

    # ---- Confirmacao (s/n) ----
    if not confirmar_prosseguimento():
        print("\n  Encerrado. A conferência foi gerada; a minuta não foi solicitada.")
        return

    # ---- Etapa 2: Geracao da minuta ----
    print("\n" + "-" * 60)
    print("  Etapa 2 — Geração da Minuta")
    print("-" * 60)
    rodar_minuta(pasta)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"ERRO: {e}", file=sys.stderr)
    finally:
        time.sleep(5)
